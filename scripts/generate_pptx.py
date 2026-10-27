#!/usr/bin/env python3
"""
generate_pptx.py

Generates a clean, simple, professional 16:9 widescreen PowerPoint presentation (.pptx)
using straightforward language, zero jargon, and strictly crediting the manual ground truth to the user.
"""

import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_FILE = BASE_DIR / "Autonomous_Driving_Annotation_and_Retrieval.pptx"

# Simple, professional color scheme
COLOR_PRIMARY = RGBColor(24, 37, 63)      # Deep Navy #18253F
COLOR_ACCENT = RGBColor(37, 99, 235)      # Royal Blue #2563EB
COLOR_TEXT = RGBColor(51, 65, 85)         # Slate Gray #334155
COLOR_MUTED = RGBColor(100, 116, 139)     # Subtle Gray #64748B
COLOR_CARD_BG = RGBColor(248, 250, 252)   # Crisp Light Background #F8FAFC
COLOR_BORDER = RGBColor(226, 232, 240)    # Soft Border #E2E8F0
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_SUCCESS = RGBColor(16, 185, 129)    # Emerald Green #10B981


def apply_header(slide, title_text, step_tag="SPRINT UPDATE"):
    # Step tracker badge
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = step_tag.upper()
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_ACCENT

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.7))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_PRIMARY


def add_card(slide, left, top, width, height, title, bullet_points, border_color=COLOR_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD_BG
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    if title:
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(13.5)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY
        p0.space_after = Pt(8)

    first = True if not title else False
    for pt in bullet_points:
        raw = pt
        stripped = raw.strip()
        if not stripped:
            if not first:
                p = tf.add_paragraph()
                p.text = ""
                p.font.size = Pt(4)
                p.space_after = Pt(2)
            continue

        is_sub = raw.startswith("  ") or raw.startswith("\t")
        # Strip any existing leading bullets, dashes, or stars so bullets never duplicate
        clean = re.sub(r'^[•\-\*\s]+', '', stripped)

        p = tf.add_paragraph() if not first else tf.paragraphs[0]
        first = False

        if stripped.endswith(":") and not is_sub and len(stripped) < 45 and not stripped.startswith("I "):
            # Subheading within card (e.g. "Road & Environment:")
            p.text = clean
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = COLOR_PRIMARY
            p.space_after = Pt(2)
        elif is_sub:
            # Indented sub-item
            if re.match(r'^\d+\.', clean):
                p.text = f"     {clean}"
            else:
                p.text = f"     •   {clean}"
            p.font.size = Pt(10)
            p.font.color.rgb = COLOR_TEXT
            p.space_after = Pt(3)
        else:
            p.text = f"•   {clean}"
            p.font.size = Pt(10.5)
            p.font.color.rgb = COLOR_TEXT
            p.space_after = Pt(5)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Simple Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)

    banner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.1))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_PRIMARY
    banner.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.3), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "SPRINT SHOWCASE"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(10)

    p1 = tf1.add_paragraph()
    p1.text = "Autonomous Driving Video Annotation & Search"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "Ground Truth Creation, AI Model Testing (Gemini vs. Qwen), and Smart Scenario Retrieval"
    p2.font.size = Pt(15)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_after = Pt(20)

    p3 = tf1.add_paragraph()
    p3.text = "Presenter: [Your Name]"
    p3.font.size = Pt(12)
    p3.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # SLIDE 2: Sprint Overview
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_header(s2, "What We Did in This Sprint", "Sprint Overview")

    add_card(
        s2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "1. Benchmark & Testing",
        [
            "Created Manual Ground Truth: Carefully annotated 200 real driving videos across 4 different datasets.",
            "Tested Two Leading AI Models: Ran Google Gemini and Qwen to automatically annotate the same videos.",
            "Generated Comparison Metrics: Evaluated how closely each AI model matched our manual ground truth."
        ]
    )

    add_card(
        s2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "2. New Search Approach",
        [
            "The Problem: Searching videos using only words/text often returns irrelevant clips or misses key actions.",
            "Our Hybrid Solution: Combined normal text search with structured driving filters (weather, road type, maneuvers, actors).",
            "Result: Instant Top-5 search results with clear match reasons."
        ]
    )

    add_card(
        s2, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "3. Scale & Delivery",
        [
            "Scaled Annotation: Successfully generated complete annotations for all 1,000 NuScenes videos using Gemini.",
            "100% Quality Validation: Every single generated file passed strict schema validation with zero errors.",
            "Ready for Next Steps: Preparing to annotate the remaining datasets to reach our full 2,100 video goal."
        ],
        border_color=COLOR_SUCCESS
    )

    # =========================================================================
    # SLIDE 3: Task 1 - Manual Ground Truth Annotation
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_header(s3, "Task 1: Creating the Manual Ground Truth Baseline", "Step 1: Ground Truth")

    add_card(
        s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Why Manual Ground Truth Was Needed",
        [
            "To measure whether an AI model is accurate, we first need reliable, verified data to compare against.",
            "I manually reviewed and annotated 200 driving videos across 4 diverse datasets:",
            "  Dashcam (50 videos): Indian city traffic, sudden turns, gates, two-wheelers.",
            "  KITTI (40 videos): European streets, cyclists, marked crosswalks.",
            "  IDD (55 videos): Indian Driving Dataset with dense traffic, rickshaws, pedestrians.",
            "  NuScenes (55 videos): Multi-lane urban roads, construction zones, trucks."
        ]
    )

    add_card(
        s3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "What Each Manual Annotation Contains",
        [
            "Road & Environment:",
            "  Road type (urban street, highway, rural road).",
            "  Weather condition (clear, rain, overcast) and lighting (day, night).",
            "",
            "Driving Scenario & Action:",
            "  Plain-English summary of what happened (e.g. 'Car turns left at intersection').",
            "  Ego car maneuver (following, turning, stopping, overtaking).",
            "",
            "Surrounding Road Users (Actors):",
            "  Other vehicles and people present (cars, pedestrians, bikes, trucks, auto-rickshaws) and their positions."
        ]
    )

    # =========================================================================
    # SLIDE 4: Task 2 - AI Model Annotation (Gemini vs. Qwen)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_header(s4, "Task 2: Automatic Annotation by AI Models", "Step 2: Automated AI Annotation")

    add_card(
        s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Google Gemini Flash-Lite",
        [
            "What it is: Google's multimodal cloud AI model.",
            "How it works: We upload the full video clip directly to Gemini via API.",
            "What it sees: It watches the entire video smoothly from start to finish.",
            "Output: Outputs a structured JSON file matching our exact schema.",
            "Strengths: Very strong at understanding context, noticing small details (like pedestrians and rickshaws), and recognizing complex turns."
        ]
    )

    add_card(
        s4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "Qwen2.5-VL (Local AI Model)",
        [
            "What it is: An open-source vision-language model.",
            "How it works: Runs locally on our own GPU server (NVIDIA RTX 6000 Ada).",
            "How it sees: Extracts 8 sampled picture frames from each video.",
            "Engineering Optimization: Configured memory settings so it runs 100% inside GPU memory, reducing processing time to just 4 seconds per video.",
            "Strengths: Fast and very accurate at spotting cars ahead in the same lane."
        ]
    )

    # =========================================================================
    # SLIDE 5: Task 3 - Generating Metrics & Comparison
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_header(s5, "Task 3: Comparing AI Models Against Manual Ground Truth", "Step 3: Accuracy & Evaluation")

    # Simple comparison table
    tbl_shape = s5.shapes.add_table(7, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(3.2))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(3.8)
    tbl.columns[1].width = Inches(2.6)
    tbl.columns[2].width = Inches(2.6)
    tbl.columns[3].width = Inches(2.7)

    table_data = [
        ("Measurement", "Gemini Flash-Lite", "Qwen2.5-VL", "Comparison"),
        ("Overall Accuracy", "69.4%", "61.9%", "Gemini (+7.5% higher)"),
        ("Overall F1 Score (Quality)", "0.50", "0.38", "Gemini was more balanced"),
        ("Finding Pedestrians", "78.4%", "33.7%", "Gemini found far more people"),
        ("Finding Auto-Rickshaws", "77.1%", "8.7%", "Gemini easily detected rickshaws"),
        ("Finding Bicycles", "71.0%", "19.1%", "Gemini was much better on bikes"),
        ("Finding Cars", "91.8%", "94.7%", "Both models were very strong")
    ]

    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLOR_PRIMARY
            else:
                p.font.color.rgb = COLOR_TEXT
                if c_idx == 3 and "Gemini" in val:
                    p.font.bold = True
                    p.font.color.rgb = COLOR_ACCENT

    # Takeaway card below table
    add_card(
        s5, Inches(0.8), Inches(4.9), Inches(11.733), Inches(1.9),
        "What We Learned from the Comparison",
        [
            "Gemini is much better at complex real-world situations: It accurately identified pedestrians, rickshaws, cyclists, and varied driving actions like turns and stops.",
            "Qwen is good for simple driving: It was very reliable at detecting lead cars, but almost always assumed the car was simply 'following traffic' and missed rare road actors."
        ]
    )

    # =========================================================================
    # SLIDE 6: Task 4 - New Approach: Hybrid Search
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_header(s6, "Task 4: New Approach - Smart Search Using Text + Structured Data", "Step 4: Smart Search")

    add_card(
        s6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Why Plain Text Search Fails",
        [
            "If an engineer wants to find: 'A car turning left at a signal in heavy rain':",
            "  Searching with words alone often misses relevant videos.",
            "  Traditional search cannot check exact driving filters (like 'must be 2 lanes' or 'traffic light must be red').",
            "",
            "Our Hybrid Solution: We score videos on TWO things at once:",
            "  1. Text Match: Does the scene description match what was typed?",
            "  2. Structured Match: Do the driving facts match (weather, road, car actions, actors)?"
        ]
    )

    add_card(
        s6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
        "How the Hybrid Search Works",
        [
            "Combined Score = Text Score + Structured Filters Score",
            "",
            "Filters Checked Automatically:",
            "  Road Network: Urban, highway, intersection, traffic lights (red/green).",
            "  Weather: Rain, clear, overcast, day or night.",
            "  Car Action: Turning left/right, following, stopping, overtaking.",
            "  Surrounding Actors: Pedestrians, trucks, bikes, auto-rickshaws.",
            "",
            "Result: Returns the Top 5 best matching videos in under 50 milliseconds, with green checkmarks showing exactly why each video matched."
        ],
        border_color=COLOR_ACCENT
    )

    # =========================================================================
    # SLIDE 7: Task 5 - Scaling Up NuScenes (1,000 Videos Completed!)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_header(s7, "Task 5: Current Progress – 1,000 NuScenes Videos Completed", "Step 5: Current Status")

    add_card(
        s7, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "1,000 Videos Processed",
        [
            "Scaled up annotation using Gemini Flash-Lite.",
            "Successfully processed all 1,000 NuScenes video clips.",
            "Handled cloud rate limits smoothly over automated batches.",
            "Managed disk storage carefully so the machine never ran out of space."
        ]
    )

    add_card(
        s7, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "100% Quality Checked",
        [
            "All 1,000 JSON files were automatically tested against our formal schema.",
            "Zero Schema Errors (1,000 / 1,000 passed).",
            "Each file contains complete road, weather, maneuver, and actor information.",
            "Ready for engineering search and downstream autonomous driving testing."
        ]
    )

    add_card(
        s7, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "Single Search Catalog Ready",
        [
            "Compiled all 1,000 video annotations into one unified catalog (dataset_catalog.json).",
            "File size is compact (only ~2 MB).",
            "Tested search live: Easily found specific rainy turns and red light stops in milliseconds."
        ],
        border_color=COLOR_SUCCESS
    )

    # =========================================================================
    # SLIDE 8: Next Steps
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_header(s8, "What We Are Doing Next", "Next Steps")

    add_card(
        s8, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "1. Complete the Full Dataset",
        [
            "We have 1,000 videos done.",
            "Next: Annotate the remaining datasets (IDD, Dashcam, KITTI) to reach the complete goal of ~2,100 videos.",
            "Process in clean batches to keep local disk space free."
        ]
    )

    add_card(
        s8, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "2. Build an Easy Search Web Page",
        [
            "Create a clean visual search website.",
            "Allow users to type what they want (e.g. 'car yielding to pedestrian') and click filter checkboxes (e.g. 'rain', 'night').",
            "Show the matching video clips directly on screen."
        ]
    )

    add_card(
        s8, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "3. Use for Safety & Edge-Case Testing",
        [
            "Use our smart search engine to pull rare, risky driving events (near-misses, sudden stops, jaywalkers).",
            "Use these extracted clips to test and validate autonomous driving software."
        ],
        border_color=COLOR_ACCENT
    )

    prs.save(OUTPUT_FILE)
    print(f"✓ Simplified PowerPoint successfully updated and saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_presentation()
